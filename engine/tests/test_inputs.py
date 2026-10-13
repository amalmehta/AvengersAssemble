import base64
import os
import random
import shutil
import struct
import subprocess
import sys
import zlib
from pathlib import Path

import pytest

from suite import inputs
from suite.config import DEFAULT_SETTINGS
from suite.orchestrator import Orchestrator
from suite.providers import MockProvider
from suite.tools import ToolError, Workspace

HAS_SIPS = Path("/usr/bin/sips").exists()


def make_png(path: Path, width: int, height: int) -> None:
    """A valid RGB PNG with a gradient, written by hand (no imaging library needed)."""
    rows = b"".join(b"\x00" + b"".join(bytes((x * 255 // width, y * 255 // height, 128)) for x in range(width))
                    for y in range(height))

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(rows, 1)) + chunk(b"IEND", b""))


def make_pdf(path: Path, page_texts: list[str]) -> None:
    """A small valid PDF with one line of Helvetica text per page."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in page_texts:
        stream = f"BT /F1 14 Tf 72 720 Td ({text}) Tj ET".encode()
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream.decode()}\nendstream")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {len(objs)} 0 R "
                    "/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


@pytest.fixture
def ws(tmp_path):
    w = Workspace(tmp_path / "workspace")
    (w.root / "inputs").mkdir()
    return w


def add(ws, name, maker=None, data=None):
    path = ws.root / "inputs" / name
    if maker:
        maker(path)
    else:
        path.write_bytes(data)
    return path


def by_name(manifests):
    return {m["name"]: m for m in manifests}


def test_pdf_short_is_native_long_is_text_with_page_ranges(ws):
    add(ws, "brief.pdf", lambda p: make_pdf(p, ["Market brief: laundry robots", "Six hours a week"]))
    add(ws, "report.pdf", lambda p: make_pdf(p, [f"Battery test page {i}" for i in range(1, 26)]))
    m = by_name(inputs.prepare_all(ws.root))
    assert m["brief.pdf"]["pages"] == 2 and "laundry robots" in m["brief.pdf"]["excerpt"]
    assert m["report.pdf"]["pages"] == 25

    short = ws.run("read_file", {"path": "inputs/brief.pdf"})
    assert short[0]["text"] == "inputs/brief.pdf (PDF, 2 pages)"
    assert short[1]["type"] == "document" and short[1]["source"]["media_type"] == "application/pdf"
    assert base64.b64decode(short[1]["source"]["data"]).startswith(b"%PDF")

    long = ws.run("read_file", {"path": "inputs/report.pdf"})
    assert isinstance(long, str) and "25 pages" in long and "Battery test page 20" in long
    assert "Battery test page 21" not in long

    sliced = ws.run("read_file", {"path": "inputs/report.pdf", "pages": "21-25"})
    assert sliced[0]["text"] == "inputs/report.pdf (PDF, pages 21-25 of 25)"
    from pypdf import PdfReader
    import io
    pdf = PdfReader(io.BytesIO(base64.b64decode(sliced[1]["source"]["data"])))
    assert len(pdf.pages) == 5 and "page 21" in pdf.pages[0].extract_text()

    for bad in ("0", "3-1", "x", "1-26"):
        with pytest.raises(ToolError):
            ws.run("read_file", {"path": "inputs/report.pdf", "pages": bad})


@pytest.mark.skipif(not HAS_SIPS, reason="needs macOS sips")
def test_images_are_downscaled_and_heic_converted(ws):
    add(ws, "small.png", lambda p: make_png(p, 300, 200))
    big = add(ws, "big.png", lambda p: make_png(p, 3000, 1000))
    subprocess.run(["/usr/bin/sips", "-s", "format", "heic", str(big), "--out", str(ws.root / "inputs" / "photo.heic")],
                   capture_output=True, check=True)
    subprocess.run(["/usr/bin/sips", "-s", "format", "tiff", str(ws.root / "inputs" / "small.png"),
                    "--out", str(ws.root / "inputs" / "scan.tiff")], capture_output=True, check=True)
    m = by_name(inputs.prepare_all(ws.root))
    assert m["scan.tiff"]["media_type"] == "image/png" and m["scan.tiff"]["width"] == 300
    assert (m["small.png"]["width"], m["small.png"]["height"]) == (300, 200)
    assert m["big.png"]["width"] == inputs.MAX_IMAGE_EDGE and m["big.png"]["original_width"] == 3000
    assert m["photo.heic"]["media_type"] == "image/jpeg" and m["photo.heic"]["width"] == inputs.MAX_IMAGE_EDGE

    result = ws.run("read_file", {"path": "inputs/photo.heic"})
    assert "downscaled from 3000×1000" in result[0]["text"]
    assert result[1]["type"] == "image" and result[1]["source"]["media_type"] == "image/jpeg"
    assert base64.b64decode(result[1]["source"]["data"])[:2] == b"\xff\xd8"  # JPEG
    assert (ws.root / "inputs" / "photo.heic").exists()  # the original stays as uploaded


def test_office_files_become_markdown(ws):
    import docx
    import openpyxl
    from pptx import Presentation

    def make_docx(p):
        d = docx.Document()
        d.add_heading("Customer interviews", 1)
        d.add_paragraph("12 of 15 parents said folding is the worst chore.")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Segment", "Pay", "Dual", "$999"
        d.save(p)

    def make_xlsx(p):
        wb = openpyxl.Workbook()
        wb.active.title = "BOM"
        wb.active.append(["Part", "Qty"])
        wb.active.append(["Gripper", 2])
        wb.save(p)

    def make_pptx(p):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Pitch"
        slide.notes_slide.notes_text_frame.text = "Lead with time saved"
        prs.save(p)

    add(ws, "interviews.docx", make_docx)
    add(ws, "bom.xlsx", make_xlsx)
    add(ws, "pitch.pptx", make_pptx)
    inputs.prepare_all(ws.root)
    text = ws.run("read_file", {"path": "inputs/interviews.docx"})
    assert "# Customer interviews" in text and "| Segment | Pay |" in text and "| Dual | $999 |" in text
    assert "## Sheet: BOM" in ws.run("read_file", {"path": "inputs/bom.xlsx"})
    pitch = ws.run("read_file", {"path": "inputs/pitch.pptx"})
    assert "## Slide 1" in pitch and "Pitch" in pitch and "*Notes:* Lead with time saved" in pitch


def test_broken_and_binary_files_fail_gracefully(ws):
    add(ws, "broken.pdf", data=b"%PDF-1.4 this is not really a pdf")
    add(ws, "blob.bin", data=bytes(range(256)))
    add(ws, "notes.txt", data="plain notes".encode())
    m = by_name(inputs.prepare_all(ws.root))
    assert m["broken.pdf"]["kind"] == "unreadable" and m["broken.pdf"]["error"]
    with pytest.raises(ToolError, match="could not be read"):
        ws.run("read_file", {"path": "inputs/broken.pdf"})
    assert "binary file" in ws.run("read_file", {"path": "inputs/blob.bin"})
    assert ws.run("read_file", {"path": "inputs/notes.txt"}) == "plain notes"
    assert "could not be read" in inputs.digest(list(m.values()))


def test_prepared_copies_are_hidden_from_listing(ws):
    add(ws, "brief.pdf", lambda p: make_pdf(p, ["hello"]))
    inputs.prepare_all(ws.root)
    assert ws.files() == ["inputs/brief.pdf"]


def _attempt(profile: str, target: Path) -> str:
    code = ("import urllib.request\n"
            f"try:\n open({str(target)!r}, 'w').write('x'); print('WRITE-OPEN')\nexcept Exception: print('write-blocked')\n"
            "try:\n urllib.request.urlopen('https://example.com', timeout=5); print('NET-OPEN')\n"
            "except Exception: print('net-blocked')\n")
    return subprocess.run(["sandbox-exec", "-p", profile, sys.executable, "-c", code],
                          capture_output=True, text=True).stdout


@pytest.mark.skipif(not shutil.which("sandbox-exec"), reason="needs macOS sandbox-exec")
def test_document_sandbox_allows_only_the_workspace(tmp_path):
    root = tmp_path / "workspace"
    root.mkdir()
    outside = tmp_path / "outside.txt"  # in the per-user temp folder: still off limits for document parsing
    out = _attempt(inputs._sandbox_profile(root), outside)
    assert "write-blocked" in out and "net-blocked" in out and not outside.exists()
    assert "WRITE-OPEN" in _attempt(inputs._sandbox_profile(root), root / "inside.txt")


@pytest.mark.skipif(not shutil.which("sandbox-exec"), reason="needs macOS sandbox-exec")
def test_image_sandbox_blocks_home_and_network(tmp_path):
    root = tmp_path / "workspace"
    root.mkdir()
    target = Path.home() / f".suite-sandbox-test-{os.getpid()}"
    try:
        out = _attempt(inputs._sandbox_profile(root, image=True), target)
        assert "write-blocked" in out and "net-blocked" in out and not target.exists()
    finally:
        target.unlink(missing_ok=True)


def test_run_prepares_inputs_and_agents_see_native_pdf(tmp_path, monkeypatch):
    monkeypatch.setenv("SUITE_HOME", str(tmp_path))
    brief = tmp_path / "brief.pdf"
    make_pdf(brief, ["Laundry robot market brief"])
    events = []
    orch = Orchestrator(MockProvider(seed=1), events.append, suite_home=tmp_path, rng=random.Random(1),
                        settings={**DEFAULT_SETTINGS, "max_parallel_tasks": 1})
    orch.run("Launch a company selling a home robot that folds laundry", [brief])
    assert [e["type"] for e in events][:2] == ["run_started", "inputs"]
    assert events[1]["files"][0] == {"name": "brief.pdf", "kind": "pdf", "bytes": brief.stat().st_size,
                                     "pages": 1, "width": None, "height": None, "error": None}
    assert "inputs/brief.pdf (PDF, 1 page," in orch.input_digest and "Laundry robot market brief" in orch.input_digest
    finished = [e for e in events if e["type"] == "task_finished"]
    assert finished and all("Read document" in e["summary"] for e in finished)
