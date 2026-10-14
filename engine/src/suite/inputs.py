"""Uploaded files → something the models can read.

At the start of a run every file in workspace/inputs/ is prepared once, into the hidden
inputs/.prepared/ folder (the original stays as uploaded):
  - images: HEIC → JPEG, TIFF and BMP → PNG; downscaled to MAX_IMAGE_EDGE px; served to agents as image blocks
  - PDFs: page count and extracted text; served as native document blocks (page ranges for long ones)
  - DOCX / XLSX / PPTX: converted to Markdown text
  - anything else readable as UTF-8: served as text

Uploads are untrusted, so all parsing happens under sandbox-exec, with no network and a time limit:
  - PDFs and Office files are parsed in a Python child process that may write only inside the workspace:
        python -m suite.inputs prepare <file> <prepared_dir>
        python -m suite.inputs slice <pdf> <first> <last> <out.pdf>
  - Images are decoded only by Apple's sips, which also needs the per-user temp folder; the engine runs
    it directly under that slightly wider profile, so no Python code parses uploads there.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

PREPARED = ".prepared"
MAX_IMAGE_EDGE = 1568           # long edge, in px: larger images are downscaled before any model sees them
MAX_IMAGE_BYTES = 3_700_000     # stays under the API's 5 MB per-image limit once base64-encoded
NATIVE_PDF_PAGES = 20           # longest PDF (or page range) sent as a native document in one read
MAX_TEXT_CHARS = 40_000         # text returned by one read of a converted file
EXCERPT_CHARS = 700             # per-file excerpt in the CEO's planning digest
CONVERT_TIMEOUT_S = 90
MAX_XLSX_ROWS = 200

IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif",
               ".webp": "image/webp"}
CONVERTED_IMAGE_TYPES = {".heic": "jpeg", ".heif": "jpeg", ".tif": "png", ".tiff": "png", ".bmp": "png"}
OFFICE_TYPES = {".docx", ".xlsx", ".pptx"}


# --- conversion (runs inside the sandboxed child process) ------------------------------

def prepare(src: Path, out_dir: Path) -> dict:
    """Make a model-ready version of one upload and describe it. Returns the manifest."""
    out_dir.mkdir(parents=True, exist_ok=True)
    ext = src.suffix.lower()
    info: dict = {"name": src.name, "bytes": src.stat().st_size}
    try:
        if ext == ".pdf":
            info.update(_prepare_pdf(src, out_dir))
        elif ext in OFFICE_TYPES:
            text = {".docx": _docx_text, ".xlsx": _xlsx_text, ".pptx": _pptx_text}[ext](src)
            info.update(_write_text(src, out_dir, text, kind="office"))
        else:
            raw = src.read_bytes()
            try:
                text = raw.decode("utf-8")
                info.update({"kind": "text", "excerpt": text[:EXCERPT_CHARS]})
            except UnicodeDecodeError:
                info.update({"kind": "binary"})
    except Exception as e:  # a broken or hostile file must not break the run
        info.update({"kind": "unreadable", "error": f"{type(e).__name__}: {e}"[:300]})
    (out_dir / f"{src.name}.json").write_text(json.dumps(info))
    return info


def _sips(*args: str, root: Path | None = None) -> subprocess.CompletedProcess:
    cmd = ["/usr/bin/sips", *args]
    if root is not None and shutil.which("sandbox-exec"):
        cmd = ["sandbox-exec", "-p", _sandbox_profile(root, image=True), *cmd]
    return subprocess.run(cmd, capture_output=True, text=True, timeout=60)


def _image_size(path: Path, root: Path) -> tuple[int, int]:
    out = _sips("-g", "pixelWidth", "-g", "pixelHeight", str(path), root=root).stdout
    nums = [int(line.split()[-1]) for line in out.splitlines() if line.strip().startswith("pixel")]
    return (nums[0], nums[1]) if len(nums) == 2 else (0, 0)


def prepare_image(src: Path, out_dir: Path, root: Path) -> dict:
    """Runs in the engine process; every decode is done by sips under the image sandbox."""
    out_dir.mkdir(parents=True, exist_ok=True)
    info: dict = {"name": src.name, "bytes": src.stat().st_size}
    try:
        info.update(_prepare_image(src, out_dir, root))
    except Exception as e:
        info.update({"kind": "unreadable", "error": f"{type(e).__name__}: {e}"[:300]})
    (out_dir / f"{src.name}.json").write_text(json.dumps(info))
    return info


def _prepare_image(src: Path, out_dir: Path, root: Path) -> dict:
    ext = src.suffix.lower()
    if not Path("/usr/bin/sips").exists():  # not on a Mac: serve the original if the API accepts it
        if ext in CONVERTED_IMAGE_TYPES:
            raise RuntimeError(f"{ext} images need macOS sips to convert")
        return {"kind": "image", "media_type": IMAGE_TYPES[ext], "prepared": src.name, "original": True}
    width, height = _image_size(src, root)
    if not width:
        raise RuntimeError("sips could not read this image")
    # The API takes PNG, JPEG, GIF and WebP; sips can't write WebP, so a WebP that needs resizing becomes JPEG.
    fmt = CONVERTED_IMAGE_TYPES.get(ext) or ("jpeg" if ext == ".webp" and max(width, height) > MAX_IMAGE_EDGE else None)
    target_ext = {"jpeg": ".jpg", "png": ".png", None: ext}[fmt]
    dest = out_dir / f"{src.name}{target_ext}"
    args = []
    if fmt == "jpeg":
        args += ["-s", "format", "jpeg", "-s", "formatOptions", "85"]
    elif fmt == "png":
        args += ["-s", "format", "png"]
    if max(width, height) > MAX_IMAGE_EDGE:
        args += ["-Z", str(MAX_IMAGE_EDGE)]
    if args:
        result = _sips(*args, str(src), "--out", str(dest), root=root)
        if not dest.exists():
            raise RuntimeError(f"sips could not convert the image: {result.stderr.strip()[:200]}")
    else:
        shutil.copy(src, dest)
    if dest.stat().st_size > MAX_IMAGE_BYTES:  # still too heavy: re-encode as JPEG
        jpg = out_dir / f"{src.name}.jpg"
        _sips("-s", "format", "jpeg", "-s", "formatOptions", "75", str(dest), "--out", str(jpg), root=root)
        if jpg != dest:
            dest.unlink(missing_ok=True)
        dest, target_ext = jpg, ".jpg"
    w, h = _image_size(dest, root)
    return {"kind": "image", "media_type": IMAGE_TYPES.get(target_ext, "image/jpeg"), "prepared": dest.name,
            "width": w, "height": h, "original_width": width, "original_height": height}


def _prepare_pdf(src: Path, out_dir: Path) -> dict:
    from pypdf import PdfReader
    reader = PdfReader(str(src))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        try:
            pages.append(f"--- Page {i} ---\n{(page.extract_text() or '').strip()}")
        except Exception:
            pages.append(f"--- Page {i} ---\n(text could not be extracted)")
    text = "\n\n".join(pages)
    has_text = any(len(p.split("\n", 1)[-1].strip()) > 20 for p in pages)
    info = _write_text(src, out_dir, text, kind="pdf")
    info.update({"pages": len(reader.pages), "has_text": has_text, "media_type": "application/pdf"})
    if not has_text:
        info["excerpt"] = "(no text layer: probably scanned; open it natively to read it)"
    return info


def _write_text(src: Path, out_dir: Path, text: str, kind: str) -> dict:
    name = f"{src.name}.md"
    (out_dir / name).write_text(text)
    return {"kind": kind, "text": name, "chars": len(text), "excerpt": text[:EXCERPT_CHARS]}


def _docx_text(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    parts = []
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        level = next((int(c) for c in style if c.isdigit()), 0) if style.startswith("heading") else 0
        parts.append(("#" * min(level, 6) + " " if level else "") + p.text)
    for n, table in enumerate(d.tables, 1):
        rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in table.rows]
        parts.append(f"\n**Table {n}**\n" + _markdown_table(rows))
    return "\n\n".join(parts)


def _xlsx_text(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            if r >= MAX_XLSX_ROWS:
                rows.append([f"… more rows not shown (sheet has {ws.max_row})"])
                break
            if any(v is not None for v in row):
                rows.append(["" if v is None else str(v) for v in row])
        parts.append(f"## Sheet: {ws.title}\n" + (_markdown_table(rows) if rows else "(empty)"))
    return "\n\n".join(parts)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        for s in slide.shapes:
            if s.has_table:
                texts.append(_markdown_table([[c.text for c in r.cells] for r in s.table.rows]))
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        parts.append(f"## Slide {i}\n" + "\n".join(texts) + (f"\n\n*Notes:* {notes}" if notes else ""))
    return "\n\n".join(parts)


def _markdown_table(rows: list[list[str]]) -> str:
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    cell = lambda v: str(v).replace("|", "\\|")
    lines = ["| " + " | ".join(cell(v) for v in rows[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(cell(v) for v in r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def slice_pdf(src: Path, first: int, last: int, out: Path) -> None:
    from pypdf import PdfReader, PdfWriter
    reader, writer = PdfReader(str(src)), PdfWriter()
    for i in range(first - 1, last):
        writer.add_page(reader.pages[i])
    with open(out, "wb") as f:
        writer.write(f)


# --- running conversion in the sandbox (engine side) -------------------------------------

def _sandbox_profile(root: Path, image: bool = False) -> str:
    """No network; writes only inside the workspace (plus the per-user temp folder for sips)."""
    extra = f' (subpath "{os.path.realpath(_user_temp_dir())}")' if image else ""
    return f"""(version 1)
(allow default)
(deny network*)
(deny file-write*)
(allow file-write* (subpath "{os.path.realpath(root)}"){extra} (literal "/dev/null") (literal "/dev/dtracehelper"))
"""


def _user_temp_dir() -> str:
    try:  # the per-user temp folder; sips writes there while converting
        return subprocess.run(["getconf", "DARWIN_USER_TEMP_DIR"], capture_output=True, text=True).stdout.strip() \
            or os.environ.get("TMPDIR", "/tmp")
    except OSError:
        return os.environ.get("TMPDIR", "/tmp")


def run_sandboxed(root: Path, *args: str) -> subprocess.CompletedProcess:
    cmd = [sys.executable, "-m", "suite.inputs", *args]
    if shutil.which("sandbox-exec"):
        cmd = ["sandbox-exec", "-p", _sandbox_profile(root), *cmd]
    env = {"PATH": "/usr/bin:/bin", "HOME": str(root), "PYTHONPATH": os.pathsep.join(sys.path)}
    return subprocess.run(cmd, cwd=root, env=env, capture_output=True, text=True, timeout=CONVERT_TIMEOUT_S)


def prepare_all(workspace_root: Path) -> list[dict]:
    """Prepare every upload in workspace/inputs/. Returns their manifests, in name order."""
    inputs = workspace_root / "inputs"
    if not inputs.is_dir():
        return []
    out_dir = inputs / PREPARED

    def one(src: Path) -> dict:
        if src.suffix.lower() in IMAGE_TYPES or src.suffix.lower() in CONVERTED_IMAGE_TYPES:
            return prepare_image(src, out_dir, workspace_root)
        try:
            result = run_sandboxed(workspace_root, "prepare", str(src), str(out_dir))
            manifest_path = out_dir / f"{src.name}.json"
            if manifest_path.exists():
                return json.loads(manifest_path.read_text())
            error = (result.stderr.strip().splitlines() or ["conversion failed"])[-1]
        except subprocess.TimeoutExpired:
            error = f"conversion took longer than {CONVERT_TIMEOUT_S} s"
        return {"name": src.name, "bytes": src.stat().st_size, "kind": "unreadable", "error": error[:300]}

    files = sorted(p for p in inputs.iterdir() if p.is_file() and not p.name.startswith("."))
    with ThreadPoolExecutor(max_workers=4) as pool:  # each file converts in its own sandboxed process
        return list(pool.map(one, files))


def manifest(workspace_root: Path, name: str) -> dict | None:
    path = workspace_root / "inputs" / PREPARED / f"{name}.json"
    return json.loads(path.read_text()) if path.exists() else None


def digest(manifests: list[dict], budget: int = 6000) -> str:
    """A short description of every upload for the CEO's planning prompt."""
    if not manifests:
        return "none"
    lines = []
    for m in manifests:
        kind = m.get("kind")
        size = f"{m.get('bytes', 0) / 1024:.0f} KB"
        if kind == "image":
            head = f"image, {m.get('width')}×{m.get('height')} px"
        elif kind == "pdf":
            head = f"PDF, {m.get('pages')} page{'s' if m.get('pages') != 1 else ''}"
        elif kind == "office":
            head = f"{Path(m['name']).suffix.lstrip('.').upper()} converted to text, {m.get('chars', 0):,} chars"
        elif kind == "unreadable":
            head = f"could not be read ({m.get('error', 'unknown error')})"
        else:
            head = kind or "file"
        line = f"- inputs/{m['name']} ({head}, {size})"
        if m.get("excerpt"):
            excerpt = " ".join(m["excerpt"].split())[:EXCERPT_CHARS]
            line += f"\n  Excerpt: {excerpt}"
        lines.append(line)
    text = "\n".join(lines)
    return text if len(text) <= budget else text[:budget] + "\n  …(digest truncated)"


def main(argv: list[str]) -> int:
    if argv[:1] == ["prepare"] and len(argv) == 3:
        prepare(Path(argv[1]), Path(argv[2]))
        return 0
    if argv[:1] == ["slice"] and len(argv) == 5:
        slice_pdf(Path(argv[1]), int(argv[2]), int(argv[3]), Path(argv[4]))
        return 0
    print("usage: python -m suite.inputs prepare <file> <dir> | slice <pdf> <first> <last> <out>", file=sys.stderr)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
